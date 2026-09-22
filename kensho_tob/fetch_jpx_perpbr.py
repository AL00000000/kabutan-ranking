# -*- coding: utf-8 -*-
"""JPXの「規模別・業種別PER・PBR」(月次)を落として、市場×33業種の平均PBRを作る。

TOB対象のPBRが低いかどうかは、**同じ時期の同じ業種と比べて**はじめて言える。
その物差しをここで作る。セルが "=1.3" のような式の形で入っているので '=' を剥がす。

  py fetch_jpx_perpbr.py
出力: jpx_perpbr.json  {"2024-01": {"プライム市場": {"7 化学": {"n":.., "pbr":..}}}}
"""
import io
import json
import re
import time
import urllib.request
from pathlib import Path

import openpyxl

BASE = Path(__file__).resolve().parent
CACHE = BASE / "raw_jpx"
OUT = BASE / "jpx_perpbr.json"
INDEX = "https://www.jpx.co.jp/markets/statistics-equities/misc/04.html"
UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) kabutan-ranking/tob-study"


def http(url, timeout=90):
    return urllib.request.urlopen(
        urllib.request.Request(url, headers={"User-Agent": UA}), timeout=timeout).read()


def cell(v):
    if v is None:
        return None
    s = str(v).lstrip("=").replace(",", "")
    try:
        return float(s)
    except ValueError:
        return None


def parse(b):
    wb = openpyxl.load_workbook(io.BytesIO(b), data_only=False)
    ws = wb["規模別・業種別（連結）"]
    out = {}
    for r in range(5, ws.max_row + 1):
        ym = ws.cell(r, 1).value
        mkt = ws.cell(r, 2).value
        ind = ws.cell(r, 4).value
        if not ym or not mkt or not ind:
            continue
        ind = re.sub(r"[　\s]+", " ", str(ind)).strip()
        n = cell(ws.cell(r, 6).value)
        pbr = cell(ws.cell(r, 8).value)
        per = cell(ws.cell(r, 7).value)
        out.setdefault(str(mkt), {})[ind] = {"n": n, "pbr": pbr, "per": per}
    return out


def main():
    CACHE.mkdir(exist_ok=True)
    h = http(INDEX).decode("utf-8", "replace")
    links = sorted(set(re.findall(r'href="([^"]+perpbr(\d{6})\.xlsx)"', h)))
    res = json.loads(OUT.read_text("utf-8")) if OUT.exists() else {}
    for path, ym in links:
        key = f"{ym[:4]}-{ym[4:]}"
        if key in res:
            continue
        f = CACHE / f"perpbr{ym}.xlsx"
        if f.exists():
            b = f.read_bytes()
        else:
            b = http("https://www.jpx.co.jp" + path)
            f.write_bytes(b)
            time.sleep(0.5)
        try:
            res[key] = parse(b)
            print(key, "ok", sum(len(v) for v in res[key].values()), "行", flush=True)
        except Exception as e:
            print(key, "FAIL", type(e).__name__, e, flush=True)
    OUT.write_text(json.dumps(res, ensure_ascii=False, indent=1), "utf-8")
    print(f"{len(res)}か月分 → {OUT.name}")


if __name__ == "__main__":
    main()
